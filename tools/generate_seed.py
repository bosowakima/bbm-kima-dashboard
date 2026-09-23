"""
Membuat ulang berkas seed SQL dari berkas Excel rekap.

Cara pakai:
    python tools/generate_seed.py rekap_sparepart.xlsx "ANALISIS MALLOMO" sparepart
    python tools/generate_seed.py rekap_service.xlsx "ANALISIS JAN-AGU 2026" service

Hasilnya ditulis ke folder supabase/. Cara yang lebih mudah untuk memperbarui data
adalah fitur Impor pada halaman Admin dashboard; skrip ini disediakan sebagai cadangan.

Prasyarat: pip install openpyxl
"""

import sys
from pathlib import Path

import openpyxl

BARIS_AWAL = 8  # baris pertama data pada sheet analisis
KOLOM = {"no": 2, "no_part": 3, "nama": 4, "harga": 5, "frekuensi": 6, "total": 7}


def baca(berkas: str, sheet: str):
    wb = openpyxl.load_workbook(berkas, data_only=True)
    ws = wb[sheet]
    hasil = []
    for baris in range(BARIS_AWAL, ws.max_row + 1):
        nomor = ws.cell(row=baris, column=KOLOM["no"]).value
        if not isinstance(nomor, int):
            break  # baris total di bawah tabel
        hasil.append(
            {
                "no": nomor,
                "no_part": ws.cell(row=baris, column=KOLOM["no_part"]).value or "-",
                "nama": ws.cell(row=baris, column=KOLOM["nama"]).value or "-",
                "harga": int(ws.cell(row=baris, column=KOLOM["harga"]).value or 0),
                "frekuensi": int(ws.cell(row=baris, column=KOLOM["frekuensi"]).value or 0),
                "total": int(ws.cell(row=baris, column=KOLOM["total"]).value or 0),
            }
        )
    return hasil


def kutip(teks) -> str:
    return str(teks).replace("'", "''")


def main() -> int:
    if len(sys.argv) != 4:
        print(__doc__)
        return 1

    berkas, sheet, departemen = sys.argv[1], sys.argv[2], sys.argv[3]
    if departemen not in {"sparepart", "service"}:
        print("Departemen harus 'sparepart' atau 'service'.")
        return 1

    data = baca(berkas, sheet)
    if not data:
        print("Tidak ada baris data yang terbaca. Periksa nama sheet.")
        return 1

    nama_keluaran = (
        "02_seed_sparepart.sql" if departemen == "sparepart" else "03_seed_service.sql"
    )
    tujuan = Path(__file__).resolve().parent.parent / "supabase" / nama_keluaran

    nilai = [
        "  ('{d}', {no}, '{part}', '{nama}', {harga}, {frek}, {total})".format(
            d=departemen,
            no=b["no"],
            part=kutip(b["no_part"]),
            nama=kutip(b["nama"]),
            harga=b["harga"],
            frek=b["frekuensi"],
            total=b["total"],
        )
        for b in data
    ]

    isi = "\n".join(
        [
            f"-- Data {departemen}: {len(data)} baris, dibuat ulang dari {Path(berkas).name}.",
            f"delete from public.rekap_barang where departemen = '{departemen}';",
            "insert into public.rekap_barang (departemen, no_urut, no_part, nama_barang, harga_satuan, frekuensi, frekuensi_total) values",
            ",\n".join(nilai) + ";",
            "",
            f"select public.rapikan_departemen('{departemen}');",
            "",
        ]
    )
    tujuan.write_text(isi, encoding="utf-8")
    print(f"{len(data)} baris ditulis ke {tujuan}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
